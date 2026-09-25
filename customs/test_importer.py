import io
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .importer import ImportFileError, build_template, import_workbook
from .models import ClassificationRule, Confidence, Garment, TariffLine

GARMENT_HEADER = ["Référence", "Nom", "Type", "Composition", "Accessoires", "Pays"]
LINE_HEADER = [
    "Pays (FR / US / CN)", "Code douanier", "Description", "Droit ad valorem (%)", "Droit spécifique",
    "Condition : fibre ou famille", "Seuil (%)", "Source", "Confiance",
]
RULE_HEADER = ["Type de vêtement", "Fibre prédominante", "Code HS (6 chiffres)", "Source", "Confiance"]


def make_workbook(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)
    return workbook


def to_file(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def report_for(reports, name):
    return next(r for r in reports if r.name == name)


class GarmentImportTests(TestCase):
    fixtures = ["customs_starter"]

    def run_import(self, rows):
        return report_for(import_workbook(to_file(make_workbook({"Vêtements": [GARMENT_HEADER, *rows]}))), "Vêtements")

    def test_creates_valid_rows_and_reports_bad_ones_with_excel_row_numbers(self):
        report = self.run_import([
            ["R1", "Chemisier A", "Chemisier", "60% coton, 40% lin", "boutons bois", "France"],
            ["R2", "T-shirt B", "T-shirt", "100% coton", "", ""],
            ["R3", "Total faux", "Chemisier", "60% coton, 30% lin", "", ""],
            ["R4", "Type inconnu", "Manteau", "100% laine", "", ""],
            ["R5", "Fibre inconnue", "Chemisier", "100% dentelle", "", ""],
        ])
        self.assertEqual(report.created, 2)
        self.assertEqual([row for row, _ in report.errors], [4, 5, 6])
        self.assertIn("100 %", report.errors[0][1])
        self.assertIn("Types disponibles", report.errors[1][1])
        self.assertEqual(report.errors[1][1].count("T-shirt, débardeur"), 1)
        self.assertIn("Fibre inconnue", report.errors[2][1])

        blouse = Garment.objects.get(reference="R1")
        self.assertEqual(blouse.garment_type, "Chemisier, chemise femme/fille (tissé)")
        self.assertEqual(blouse.accessory_materials, ["bois"])
        self.assertEqual(Garment.objects.get(reference="R2").origin_country, "France")

    def test_same_reference_updates_instead_of_duplicating(self):
        self.run_import([["R1", "Chemisier A", "Chemisier", "100% coton", "", ""]])
        report = self.run_import([["R1", "Chemisier A", "Chemisier", "70% coton, 30% lin", "", "Portugal"]])
        self.assertEqual((report.created, report.updated), (0, 1))
        garment = Garment.objects.get()
        self.assertEqual(garment.fibres.count(), 2)
        self.assertEqual(garment.origin_country, "Portugal")

    def test_accessory_keywords(self):
        self.run_import([["R1", "A", "Chemisier", "100% coton", "boutons en nacre, zip", ""]])
        self.assertEqual(Garment.objects.get().accessory_materials, ["corne_nacre", "metal"])

    def test_missing_required_column(self):
        workbook = make_workbook({"Vêtements": [["Référence", "Nom"], ["R1", "A"]]})
        report = report_for(import_workbook(to_file(workbook)), "Vêtements")
        self.assertIn("Colonne(s) obligatoire(s)", report.errors[0][1])

    def test_blank_rows_are_ignored(self):
        report = self.run_import([[None] * 6, ["R1", "A", "Chemisier", "100% coton", "", ""], [None] * 6])
        self.assertEqual((report.created, len(report.errors)), (1, 0))


class LineImportTests(TestCase):
    fixtures = ["customs_starter"]

    def run_import(self, rows, staff=True):
        workbook = make_workbook({"Lignes tarifaires": [LINE_HEADER, *rows]})
        return report_for(import_workbook(to_file(workbook), allow_reference_data=staff), "Lignes tarifaires")

    def test_creates_lines_with_various_rate_formats(self):
        report = self.run_import([
            ["FR", "6206 30 00", "Chemisiers de coton", 12, "", "", None, "TARIC", "à vérifier"],
            ["CN", "6206.30.00", "Blouses de coton", "16,5 %", "", "", None, "", ""],
            ["US", "6206.30.99", "Test lin", "3,5", "", "lin", 36, "", "brouillon"],
            ["FR", "6205.20.00", "Chemises hommes coton", "Free", "", "", None, "", ""],
        ])
        self.assertEqual((report.created, len(report.errors)), (4, 0))
        fr = TariffLine.objects.get(jurisdiction="FR", code="6206 30 00")
        self.assertEqual((fr.hs6, fr.ad_valorem_rate, fr.confidence), ("620630", Decimal("0.12"), Confidence.TO_CHECK))
        self.assertEqual(TariffLine.objects.get(jurisdiction="CN").ad_valorem_rate, Decimal("0.165"))
        lin = TariffLine.objects.get(code="6206.30.99")
        self.assertEqual((lin.condition_fibre, lin.condition_min_pct), ("lin", Decimal("36")))
        self.assertEqual(TariffLine.objects.get(code="6205.20.00").ad_valorem_rate, Decimal("0"))

    def test_excel_percent_formatted_cell(self):
        workbook = make_workbook({"Lignes tarifaires": [LINE_HEADER, ["CN", "6109.10.00", "T-shirts coton", 0.14]]})
        workbook["Lignes tarifaires"]["D2"].number_format = "0.0%"
        report = report_for(import_workbook(to_file(workbook), allow_reference_data=True), "Lignes tarifaires")
        self.assertEqual(report.created, 1)
        self.assertEqual(TariffLine.objects.get(jurisdiction="CN").ad_valorem_rate, Decimal("0.14"))

    def test_bad_rows_are_reported_not_saved(self):
        report = self.run_import([
            ["JP", "1234.56.00", "Japon", 5],
            ["FR", 6206.3, "Code devenu nombre", 5],
            ["FR", "6206.30.00", "Validé", 5, "", "", None, "", "validé par l'experte"],
            ["FR", "6206.30.00", "Condition sans seuil", 5, "", "lin", None],
            ["FR", "6206.30.00", "Trop de décimales", "0,165"],
        ])
        self.assertEqual(report.created, 0)
        self.assertEqual(len(report.errors), 5)
        self.assertIn("Texte", report.errors[1][1])
        self.assertIn("ne s'importe pas", report.errors[2][1])
        self.assertFalse(TariffLine.objects.filter(jurisdiction="FR").exists())

    def test_changing_a_validated_line_removes_validation(self):
        line = TariffLine.objects.get(code="6206.30.30")
        line.validated_by_expert = True
        line.confidence = Confidence.VALIDATED
        line.save()

        report = self.run_import([["US", "6206.30.30", line.description, 20]])
        self.assertEqual(report.updated, 1)
        self.assertIn("validation", report.warnings[0][1])
        line.refresh_from_db()
        self.assertFalse(line.validated_by_expert)
        self.assertEqual(line.confidence, Confidence.TO_CHECK)
        self.assertEqual(line.ad_valorem_rate, Decimal("0.2"))
        self.assertIsNone(line.last_verified)

    def test_identical_line_is_unchanged_and_keeps_validation(self):
        line = TariffLine.objects.get(code="6206.30.30")
        line.validated_by_expert = True
        line.save()
        report = self.run_import([["US", "6206.30.30", line.description, "15,4"]])
        self.assertEqual((report.unchanged, report.updated), (1, 0))
        line.refresh_from_db()
        self.assertTrue(line.validated_by_expert)

    def test_code_written_differently_matches_existing_line(self):
        report = self.run_import([["US", "62063030", "Chemisiers coton", "15,4"]])
        self.assertEqual(report.updated, 1)
        self.assertEqual(TariffLine.objects.filter(jurisdiction="US", hs6="620630").count(), 3)

    def test_non_staff_cannot_import_reference_data(self):
        report = self.run_import([["FR", "6206 30 00", "X", 12]], staff=False)
        self.assertTrue(report.skipped_reason)
        self.assertFalse(TariffLine.objects.filter(jurisdiction="FR").exists())


class RuleImportTests(TestCase):
    fixtures = ["customs_starter"]

    def run_import(self, rows):
        workbook = make_workbook({"Règles de classement": [RULE_HEADER, *rows]})
        return report_for(import_workbook(to_file(workbook), allow_reference_data=True), "Règles de classement")

    def test_new_type_and_families(self):
        report = self.run_import([
            ["Chemise homme (tissé)", "Coton", 620520, "", ""],
            ["Chemise homme (tissé)", "Fibres synthétiques", "6205.30", "", ""],
            ["Chemise homme (tissé)", "Bizarre", "6205.90", "", ""],
            ["Chemise homme (tissé)", "Soie", 6205.9, "", ""],
        ])
        self.assertEqual((report.created, len(report.errors)), (2, 2))
        self.assertEqual(ClassificationRule.objects.get(garment_type="Chemise homme (tissé)", fibre_group="coton").hs6, "620520")
        self.assertIn("Texte", report.errors[1][1])

    def test_existing_rule_changed_loses_validation(self):
        rule = ClassificationRule.objects.get(garment_type__startswith="Chemisier", fibre_group="coton")
        rule.validated_by_expert = True
        rule.save()
        report = self.run_import([[rule.garment_type, "coton", "620699", "", ""]])
        self.assertEqual(report.updated, 1)
        rule.refresh_from_db()
        self.assertEqual(rule.hs6, "620699")
        self.assertFalse(rule.validated_by_expert)


class WorkbookLevelTests(TestCase):
    fixtures = ["customs_starter"]

    def test_unreadable_file(self):
        with self.assertRaises(ImportFileError):
            import_workbook(io.BytesIO(b"pas un classeur excel"))

    def test_no_known_sheet(self):
        with self.assertRaises(ImportFileError):
            import_workbook(to_file(make_workbook({"Feuil1": [["a"]]})))

    def test_template_is_valid_and_imports_cleanly_when_empty(self):
        content = build_template()
        names = load_workbook(io.BytesIO(content)).sheetnames
        self.assertEqual(names, ["Mode d'emploi", "Vêtements", "Lignes tarifaires", "Règles de classement", "Listes"])
        reports = import_workbook(io.BytesIO(content), allow_reference_data=True)
        self.assertEqual(len(reports), 3)
        for report in reports:
            self.assertEqual((report.created, report.updated, report.errors), (0, 0, []))


class ImportViewTests(TestCase):
    fixtures = ["customs_starter"]

    def login(self, staff=False):
        user = get_user_model().objects.create_user("u", password="not-a-real-password", is_staff=staff)
        self.client.force_login(user)

    def upload(self, buffer, name="import.xlsx"):
        return self.client.post(reverse("garment_import"), {"file": SimpleUploadedFile(name, buffer.getvalue())})

    def test_requires_login(self):
        self.assertEqual(self.client.get(reverse("garment_import")).status_code, 302)
        self.assertEqual(self.client.get(reverse("garment_import_template")).status_code, 302)

    def test_upload_creates_garments_and_shows_report(self):
        self.login()
        workbook = make_workbook({"Vêtements": [GARMENT_HEADER, ["R1", "Chemisier A", "Chemisier", "100% coton", "", ""]]})
        response = self.upload(to_file(workbook))
        self.assertContains(response, "1 créé")
        self.assertEqual(Garment.objects.count(), 1)

    def test_rejects_non_xlsx_and_garbage(self):
        self.login()
        self.assertContains(self.upload(io.BytesIO(b"x"), name="import.csv"), "format .xlsx")
        self.assertContains(self.upload(io.BytesIO(b"pas un classeur"), name="import.xlsx"), "Fichier illisible")

    def test_reference_sheets_only_for_staff(self):
        workbook = make_workbook({"Lignes tarifaires": [LINE_HEADER, ["FR", "6206 30 00", "X", 12]]})
        self.login(staff=False)
        self.assertContains(self.upload(to_file(workbook)), "réservée aux comptes équipe")
        self.assertFalse(TariffLine.objects.filter(jurisdiction="FR").exists())

    def test_template_download(self):
        self.login()
        response = self.client.get(reverse("garment_import_template"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
