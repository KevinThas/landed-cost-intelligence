import io
import zipfile
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from types import SimpleNamespace

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import quote_sheetname
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

from .fibres import (
    FIBRE_CHOICES, GROUP_CHOICES, composition_total_error, normalize_text, parse_composition, resolve_fibre,
)
from .models import (
    ACCESSORY_CHOICES, ClassificationRule, Confidence, Garment, GarmentFibre, TariffLine, digits_only,
)

MAX_ROWS = 5000
MAX_SCAN_ROWS = 20000

SHEET_GARMENTS = "Vêtements"
SHEET_LINES = "Lignes tarifaires"
SHEET_RULES = "Règles de classement"
SHEET_GUIDE = "Mode d'emploi"
SHEET_LISTS = "Listes"

# champ : (libellé de l'en-tête dans le modèle, autres intitulés acceptés)
GARMENT_COLUMNS = {
    "reference": ("Référence", ["ref", "sku", "code article"]),
    "name": ("Nom du vêtement", ["nom", "designation", "produit", "libelle"]),
    "garment_type": ("Type de vêtement", ["type", "categorie", "famille"]),
    "composition": ("Composition", ["matiere", "matieres", "composition du tissu principal"]),
    "accessories": ("Accessoires non textiles", ["accessoires"]),
    "origin_country": ("Pays de fabrication", ["pays", "origine", "pays d origine"]),
    "notes": ("Notes", ["commentaires", "remarques"]),
}
LINE_COLUMNS = {
    "jurisdiction": ("Pays (FR / US / CN)", ["pays", "zone", "pays zone", "juridiction"]),
    "code": ("Code douanier", ["code", "code national", "nomenclature"]),
    "description": ("Description", ["designation", "libelle"]),
    "rate": ("Droit ad valorem (%)", ["droit", "taux", "taux %", "droit de douane"]),
    "specific": ("Droit spécifique", ["specifique"]),
    "condition_fibre": ("Condition : fibre ou famille", ["condition fibre", "condition de fibre"]),
    "condition_min": ("Seuil (%)", ["seuil", "seuil minimum", "condition seuil"]),
    "source": ("Source", []),
    "confidence": ("Confiance", ["niveau de confiance"]),
    "notes": ("Notes", ["commentaires", "remarques"]),
}
RULE_COLUMNS = {
    "garment_type": ("Type de vêtement", ["type"]),
    "fibre_group": ("Fibre prédominante", ["fibre", "famille de fibres", "famille"]),
    "hs6": ("Code HS (6 chiffres)", ["code hs", "hs6", "code harmonise", "hs"]),
    "source": ("Source", []),
    "confidence": ("Confiance", ["niveau de confiance"]),
    "notes": ("Notes", ["commentaires", "remarques"]),
}

_BLANK = SimpleNamespace(value=None, number_format="General")


class ImportFileError(Exception):
    """Le fichier est inexploitable (illisible, aucune feuille reconnue)."""


@dataclass
class SheetReport:
    name: str
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)
    skipped_reason: str = ""


# ---------------------------------------------------------------- lecture des cellules

def _is_blank(cell):
    value = cell.value
    return value is None or (isinstance(value, str) and not value.strip())


def text(cell):
    if _is_blank(cell):
        return ""
    value = cell.value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def code_text(cell, label):
    value = cell.value
    if isinstance(value, float) and not value.is_integer():
        raise ValueError(
            f"{label} : Excel a transformé le code en nombre ({value}). "
            "Mettez la colonne au format Texte, puis ressaisissez le code."
        )
    return text(cell)


def percent_points(cell):
    """Renvoie un pourcentage en points (16,5 pour 16,5 %), ou None si la cellule est vide."""
    value = cell.value
    if _is_blank(cell):
        return None
    if isinstance(value, str):
        raw = normalize_text(value)
        if raw in {"free", "franc", "exonere", "aucun"}:
            return Decimal(0)
        try:
            points = Decimal(value.strip().replace("%", "").replace(",", ".").replace(" ", "").replace(" ", ""))
        except InvalidOperation:
            raise ValueError(f"« {value} » n'est pas un pourcentage valide.")
    elif isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError(f"« {value} » n'est pas un pourcentage valide.")
    else:
        points = Decimal(str(value))
        if "%" in (getattr(cell, "number_format", "") or ""):
            points *= 100
    rounded = points.quantize(Decimal("0.01"))
    if abs(points - rounded) > Decimal("0.0005"):
        raise ValueError(f"« {value} » : saisissez le pourcentage en points, 2 décimales maximum (16,5 pour 16,5 %).")
    return rounded


def _fraction(points):
    if points is None:
        return None
    if not 0 <= points <= 500:
        raise ValueError(f"Un droit de {points} % est hors limites.")
    return (points / 100).quantize(Decimal("0.0001"))


def _message(error):
    return "; ".join(error.messages) if isinstance(error, ValidationError) else str(error)


def read_rows(worksheet, columns, required):
    """Retourne [(numéro de ligne Excel, {champ: cellule})]. La première ligne non vide est l'en-tête."""
    aliases = {}
    for name, (label, others) in columns.items():
        for alias in [label, *others]:
            aliases.setdefault(normalize_text(alias), name)

    mapping = None
    rows = []
    for scanned, row in enumerate(worksheet.iter_rows(), start=1):
        if scanned > MAX_SCAN_ROWS:
            break
        if all(_is_blank(c) for c in row):
            continue
        if mapping is None:
            mapping = {}
            for index, cell in enumerate(row):
                if not _is_blank(cell):
                    name = aliases.get(normalize_text(str(cell.value)))
                    if name and name not in mapping:
                        mapping[name] = index
            missing = [columns[name][0] for name in required if name not in mapping]
            if missing:
                raise ValueError("Colonne(s) obligatoire(s) introuvable(s) : " + ", ".join(missing) + ".")
            continue
        rows.append((row[0].row, {name: row[i] for name, i in mapping.items() if i < len(row)}))
    return rows


def _cell(values, name):
    return values.get(name, _BLANK)


# ---------------------------------------------------------------- résolution des valeurs

_JURISDICTIONS = {
    "fr": "FR", "france": "FR", "ue": "FR", "eu": "FR", "france ue": "FR", "union europeenne": "FR",
    "us": "US", "usa": "US", "etats unis": "US", "etats unis d amerique": "US",
    "cn": "CN", "chine": "CN", "china": "CN",
}

_GROUP_LOOKUP = {}
for _code, _label in GROUP_CHOICES:
    _GROUP_LOOKUP[normalize_text(_code)] = _code
    _GROUP_LOOKUP[normalize_text(_label)] = _code
_GROUP_LOOKUP.update({
    "poils fins": "laine", "laine ou poils fins": "laine",
    "synthetiques": "fibres_chimiques", "artificielles": "fibres_chimiques",
    "fibres synthetiques": "fibres_chimiques", "fibres chimiques": "fibres_chimiques",
    "synthetiques ou artificielles": "fibres_chimiques",
})

_CONFIDENCES = {
    "a verifier": Confidence.TO_CHECK,
    "brouillon": Confidence.DRAFT,
    "faible confiance": Confidence.LOW_CONFIDENCE,
    "brouillon faible confiance": Confidence.LOW_CONFIDENCE,
}

_ACCESSORY_WORDS = {
    "bois": "bois", "metal": "metal", "laiton": "metal", "acier": "metal", "zip": "metal",
    "plastique": "plastique", "resine": "plastique", "cuir": "cuir", "peau": "cuir",
    "corne": "corne_nacre", "nacre": "corne_nacre", "os": "corne_nacre", "verre": "verre",
}
_NO_ACCESSORY = {"aucun", "aucune", "non", "sans", "none", "0"}


def resolve_jurisdiction(value):
    code = _JURISDICTIONS.get(normalize_text(value))
    if code is None:
        raise ValueError(f"Pays / zone inconnu : « {value} ». Utilisez FR, US ou CN.")
    return code


def resolve_group(value):
    code = _GROUP_LOOKUP.get(normalize_text(value))
    if code is None:
        raise ValueError(
            f"Fibre prédominante inconnue : « {value} ». Familles : "
            + ", ".join(label.lower() for _, label in GROUP_CHOICES) + "."
        )
    return code


def resolve_condition(value):
    """Renvoie (famille, fibre) : au plus l'un des deux est rempli."""
    if not value:
        return "", ""
    key = normalize_text(value)
    if key in _GROUP_LOOKUP:
        return _GROUP_LOOKUP[key], ""
    fibre = resolve_fibre(value)
    if fibre:
        return "", fibre
    raise ValueError(f"Fibre ou famille de fibres inconnue : « {value} ».")


def resolve_confidence(value):
    if not value:
        return None
    key = normalize_text(value)
    if key.startswith("valide"):
        raise ValueError(
            "Le niveau « validé » ne s'importe pas : il se pose dans l'administration, par l'experte."
        )
    level = _CONFIDENCES.get(key)
    if level is None:
        raise ValueError(f"Niveau de confiance inconnu : « {value} » (à vérifier, brouillon, faible confiance).")
    return level


def resolve_garment_type(value, types):
    key = normalize_text(value)
    if not key:
        raise ValueError("Type de vêtement manquant.")
    if not types:
        raise ValueError("Aucune règle de classement n'existe encore : importez d'abord la feuille des règles.")
    normalized = {t: normalize_text(t) for t in types}
    exact = [t for t, n in normalized.items() if n == key]
    if exact:
        return exact[0]
    partial = [t for t, n in normalized.items() if key in n]
    if len(partial) == 1:
        return partial[0]
    raise ValueError(f"Type de vêtement « {value} » inconnu ou ambigu. Types disponibles : {'; '.join(sorted(types))}.")


def parse_accessories(value):
    key = normalize_text(value)
    if not key or key in _NO_ACCESSORY:
        return []
    found = []
    for word in key.split():
        code = _ACCESSORY_WORDS.get(word)
        if code and code not in found:
            found.append(code)
    return found or ["autre"]


# ---------------------------------------------------------------- traitement des lignes

def _reset_validation(obj, warnings, what):
    if obj.validated_by_expert or obj.confidence == Confidence.VALIDATED:
        obj.validated_by_expert = False
        if obj.confidence == Confidence.VALIDATED:
            obj.confidence = Confidence.TO_CHECK
        warnings.append(f"{what} modifiée : la validation par l'experte a été retirée (à revalider).")


def _apply_meta(obj, values):
    """Source, confiance et notes : appliquées seulement si la colonne est renseignée."""
    changed = False
    for name in ("source", "notes"):
        new = text(_cell(values, name))
        if new and new != getattr(obj, name):
            setattr(obj, name, new)
            changed = True
    confidence = resolve_confidence(text(_cell(values, "confidence")))
    if confidence and confidence != obj.confidence:
        obj.confidence = confidence
        changed = True
    return changed


def import_garment_row(values, types):
    name = text(_cell(values, "name"))
    if not name:
        raise ValueError("Nom du vêtement manquant.")
    garment_type = resolve_garment_type(text(_cell(values, "garment_type")), types)
    fibres = parse_composition(text(_cell(values, "composition")))
    error = composition_total_error(fibres)
    if error:
        raise ValueError(error)

    reference = text(_cell(values, "reference"))
    garment = Garment.objects.filter(reference=reference).first() if reference else None
    outcome = "created" if garment is None else "updated"
    if garment is None:
        garment = Garment(reference=reference)
    garment.name = name
    garment.garment_type = garment_type
    garment.origin_country = text(_cell(values, "origin_country")) or "France"
    garment.accessory_materials = parse_accessories(text(_cell(values, "accessories")))
    garment.notes = text(_cell(values, "notes"))
    garment.full_clean()
    garment.save()
    garment.fibres.all().delete()
    GarmentFibre.objects.bulk_create(GarmentFibre(garment=garment, fibre=c, percent=p) for c, p in fibres)
    return outcome, []


def import_line_row(values):
    jurisdiction = resolve_jurisdiction(text(_cell(values, "jurisdiction")))
    code = code_text(_cell(values, "code"), "Code douanier")
    digits = digits_only(code)
    if len(digits) < 6:
        raise ValueError("Le code douanier doit contenir au moins 6 chiffres.")
    description = text(_cell(values, "description"))
    if not description:
        raise ValueError("Description manquante.")

    rate = _fraction(percent_points(_cell(values, "rate")))
    specific = text(_cell(values, "specific"))
    group, fibre = resolve_condition(text(_cell(values, "condition_fibre")))
    threshold = percent_points(_cell(values, "condition_min"))
    if threshold is not None and not 0 < threshold <= 100:
        raise ValueError("Le seuil doit être compris entre 0 et 100 %.")
    core = {
        "description": description, "ad_valorem_rate": rate, "specific_duty": specific,
        "condition_group": group, "condition_fibre": fibre, "condition_min_pct": threshold,
    }

    line = next(
        (l for l in TariffLine.objects.filter(jurisdiction=jurisdiction, hs6=digits[:6]) if digits_only(l.code) == digits),
        None,
    )
    warnings = []
    if line is None:
        line = TariffLine(jurisdiction=jurisdiction, code=code, hs6=digits[:6], **core)
        _apply_meta(line, values)
        outcome = "created"
    else:
        core_changed = any(getattr(line, k) != v for k, v in core.items())
        for key, new in core.items():
            setattr(line, key, new)
        meta_changed = _apply_meta(line, values)
        if core_changed:
            _reset_validation(line, warnings, "Ligne")
            line.last_verified = None
        outcome = "updated" if core_changed or meta_changed else "unchanged"

    if rate is None and not specific:
        warnings.append("Aucun droit renseigné (ni ad valorem, ni spécifique).")
    if outcome != "unchanged":
        line.full_clean()
        line.save()
    return outcome, warnings


def import_rule_row(values):
    garment_type = text(_cell(values, "garment_type"))
    if not garment_type:
        raise ValueError("Type de vêtement manquant.")
    group = resolve_group(text(_cell(values, "fibre_group")))
    hs6 = digits_only(code_text(_cell(values, "hs6"), "Code HS"))
    if len(hs6) != 6:
        raise ValueError("Le code HS comporte exactement 6 chiffres.")

    key = normalize_text(garment_type)
    existing = next(
        (r for r in ClassificationRule.objects.filter(fibre_group=group) if normalize_text(r.garment_type) == key),
        None,
    )
    warnings = []
    if existing is None:
        rule = ClassificationRule(garment_type=garment_type, fibre_group=group, hs6=hs6)
        _apply_meta(rule, values)
        outcome = "created"
    else:
        rule = existing
        hs_changed = rule.hs6 != hs6
        rule.hs6 = hs6
        meta_changed = _apply_meta(rule, values)
        if hs_changed:
            _reset_validation(rule, warnings, "Règle")
            rule.last_verified = None
        outcome = "updated" if hs_changed or meta_changed else "unchanged"
    if outcome != "unchanged":
        rule.full_clean()
        rule.save()
    return outcome, warnings


def _run_sheet(report, worksheet, columns, required, handler):
    try:
        rows = read_rows(worksheet, columns, required)
    except ValueError as error:
        report.errors.append((None, str(error)))
        return
    if len(rows) > MAX_ROWS:
        report.errors.append((None, f"Plus de {MAX_ROWS} lignes : découpez le fichier en plusieurs."))
        return
    for number, values in rows:
        try:
            with transaction.atomic():
                outcome, warnings = handler(values)
        except (ValueError, ValidationError) as error:
            report.errors.append((number, _message(error)))
            continue
        setattr(report, outcome, getattr(report, outcome) + 1)
        report.warnings.extend((number, w) for w in warnings)


def import_workbook(file, allow_reference_data=False):
    """Importe un classeur. Les lignes en erreur sont ignorées et listées, les autres sont enregistrées."""
    try:
        workbook = load_workbook(file, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as error:
        raise ImportFileError("Fichier illisible : envoyez un classeur Excel au format .xlsx.") from error

    sheets = {normalize_text(ws.title): ws for ws in workbook.worksheets}
    plan = [
        (SHEET_RULES, RULE_COLUMNS, {"garment_type", "fibre_group", "hs6"}, import_rule_row, True),
        (SHEET_LINES, LINE_COLUMNS, {"jurisdiction", "code", "description"}, import_line_row, True),
        (SHEET_GARMENTS, GARMENT_COLUMNS, {"name", "garment_type", "composition"}, None, False),
    ]
    if not any(normalize_text(title) in sheets for title, *_ in plan):
        raise ImportFileError(
            f"Aucune feuille reconnue. Feuilles attendues : {SHEET_GARMENTS}, {SHEET_LINES}, {SHEET_RULES} "
            "(téléchargez le modèle)."
        )

    reports = []
    for title, columns, required, handler, staff_only in plan:
        worksheet = sheets.get(normalize_text(title))
        if worksheet is None:
            continue
        report = SheetReport(name=title)
        reports.append(report)
        if staff_only and not allow_reference_data:
            report.skipped_reason = "Feuille ignorée : réservée aux comptes équipe (administration)."
            continue
        if handler is None:
            types = sorted(set(ClassificationRule.objects.values_list("garment_type", flat=True)))
            handler = lambda values, types=types: import_garment_row(values, types)
        _run_sheet(report, worksheet, columns, required, handler)
    return reports


# ---------------------------------------------------------------- modèle Excel à télécharger

GUIDE = [
    "MODE D'EMPLOI",
    "",
    "Remplissez une ou plusieurs feuilles, puis envoyez le fichier dans l'application (Importer).",
    "Les feuilles vides sont ignorées. Les lignes en erreur sont refusées et listées, les autres sont enregistrées.",
    "",
    "Feuille « Vêtements » : un vêtement par ligne.",
    "  - Composition : comme sur l'étiquette, ex. 60% coton, 30% soie, 10% viscose (total = 100 %).",
    "  - Accessoires : boutons, zip, patch… (bois, métal, plastique, cuir, corne, nacre, verre). Ils ne comptent pas dans les %.",
    "  - Type de vêtement : voir la feuille « Listes » (le début du nom suffit si c'est sans ambiguïté).",
    "  - Une Référence déjà connue met à jour le vêtement existant au lieu d'en créer un second.",
    "",
    "Feuille « Lignes tarifaires » : le tarif douanier de chaque pays (réservé aux comptes équipe).",
    "  - Pays : FR, US ou CN. Code : tel que publié (ex. 6206.30.30). Droit ad valorem : en %, ex. 15,4.",
    "  - Droit spécifique : ex. 56,3 ¢/kg (s'ajoute au droit ad valorem).",
    "  - Condition + Seuil : pour « contenant 36 % ou plus de lin » -> lin et 36.",
    "  - Si vous modifiez un droit déjà validé, la validation est retirée (à revalider).",
    "",
    "Feuille « Règles de classement » : quel code HS (6 chiffres) pour un type de vêtement et une fibre prédominante.",
    "",
    "Le niveau « validé par l'experte » ne s'importe pas : il se pose dans l'administration.",
    "Confiance acceptée : à vérifier, brouillon, faible confiance. Si vous laissez vide : brouillon pour une nouvelle ligne.",
    "Format des codes : gardez la colonne au format Texte pour qu'Excel ne supprime pas les zéros et les points.",
]


def build_template():
    workbook = Workbook()
    guide = workbook.active
    guide.title = SHEET_GUIDE
    for line in GUIDE:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    guide.column_dimensions["A"].width = 120

    def data_sheet(title, columns, text_columns, widths):
        sheet = workbook.create_sheet(title)
        sheet.append([label for label, _ in columns.values()])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.freeze_panes = "A2"
        for letter, width in zip("ABCDEFGHIJ", widths):
            sheet.column_dimensions[letter].width = width
        for letter in text_columns:
            sheet.column_dimensions[letter].number_format = "@"
        return sheet

    garments = data_sheet(SHEET_GARMENTS, GARMENT_COLUMNS, "A", [18, 32, 44, 44, 26, 22, 30])
    data_sheet(SHEET_LINES, LINE_COLUMNS, "B", [16, 16, 60, 18, 18, 26, 12, 34, 16, 30])
    data_sheet(SHEET_RULES, RULE_COLUMNS, "C", [44, 30, 20, 34, 16, 30])

    lists = workbook.create_sheet(SHEET_LISTS)
    types = sorted(set(ClassificationRule.objects.values_list("garment_type", flat=True)))
    columns = [
        ("Types de vêtement", types),
        ("Fibres", [label for _, label in FIBRE_CHOICES]),
        ("Familles de fibres", [label for _, label in GROUP_CHOICES]),
        ("Accessoires", [label for _, label in ACCESSORY_CHOICES]),
        ("Pays / zone", ["FR", "US", "CN"]),
        ("Confiance", ["à vérifier", "brouillon", "faible confiance"]),
    ]
    for index, (title, items) in enumerate(columns, start=1):
        lists.cell(row=1, column=index, value=title).font = Font(bold=True)
        for row, item in enumerate(items, start=2):
            lists.cell(row=row, column=index, value=item)
        lists.column_dimensions[lists.cell(row=1, column=index).column_letter].width = 44

    if types:
        picker = DataValidation(
            type="list", formula1=f"{quote_sheetname(SHEET_LISTS)}!$A$2:$A${len(types) + 1}",
            allow_blank=True, showErrorMessage=False,
        )
        garments.add_data_validation(picker)
        picker.add("C2:C5000")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
